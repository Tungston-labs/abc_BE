
import os
import requests

from django.core.management.base import BaseCommand
from django.conf import settings

from customers.models import Customer

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment


API_URL = f"{settings.SIGNAL_API_BASE_URL}/signals"


class Command(BaseCommand):
    help = (
        "Check MIB ONTs from Signal API against Customer database "
        "and generate an Excel report"
    )

    def handle(self, *args, **kwargs):

        try:
            # =====================================================
            # 1. Fetch Signal API
            # =====================================================

            self.stdout.write("")
            self.stdout.write("=" * 80)
            self.stdout.write("MIB ONT CHECK")
            self.stdout.write("=" * 80)

            self.stdout.write(
                f"Fetching Signal API: {API_URL}"
            )

            response = requests.get(
                API_URL,
                timeout=30
            )

            response.raise_for_status()

            data = response.json()

            # Make sure response is a list
            if not isinstance(data, list):
                self.stderr.write(
                    "❌ Signal API response is not a list."
                )
                self.stderr.write(
                    f"Received type: {type(data).__name__}"
                )
                return

            self.stdout.write(
                f"Signal API records: {len(data)}"
            )

            # =====================================================
            # 2. Extract serial numbers from Signal API
            # =====================================================

            api_serials = {
                str(item.get("serial_number")).strip()
                for item in data
                if isinstance(item, dict)
                and item.get("serial_number")
            }

            self.stdout.write(
                f"Unique MIB serial numbers: {len(api_serials)}"
            )

            # =====================================================
            # 3. Get Customer ONT numbers from Django DB
            # =====================================================

            customer_serials = set(
                Customer.objects
                .exclude(ont_number__isnull=True)
                .exclude(ont_number="")
                .values_list("ont_number", flat=True)
            )

            customer_serials = {
                str(serial).strip()
                for serial in customer_serials
                if serial
            }

            self.stdout.write(
                f"Customer ONT numbers in Django: "
                f"{len(customer_serials)}"
            )

            # =====================================================
            # 4. Compare MIB serials with Django Customer ONTs
            # =====================================================

            missing_serials = api_serials - customer_serials

            self.stdout.write(
                f"Missing ONTs in Django: "
                f"{len(missing_serials)}"
            )

            # =====================================================
            # 5. Get complete API records for missing ONTs
            # =====================================================

            missing_onts = []

            for item in data:

                if not isinstance(item, dict):
                    continue

                serial = item.get("serial_number")

                if not serial:
                    continue

                serial = str(serial).strip()

                if serial in missing_serials:
                    missing_onts.append(item)

            # =====================================================
            # 6. Create Excel workbook
            # =====================================================

            workbook = Workbook()

            worksheet = workbook.active
            worksheet.title = "Missing MIB ONTs"

            # =====================================================
            # 7. Determine all fields returned by Signal API
            # =====================================================

            all_fields = set()

            for item in missing_onts:
                all_fields.update(item.keys())

            # Put serial_number first
            field_list = []

            if "serial_number" in all_fields:
                field_list.append("serial_number")
                all_fields.remove("serial_number")

            # Prefer these fields near the beginning
            preferred_fields = [
                "olt",
                "olt_ip",
                "olt_name",
                "mac_address",
                "mac",
                "port",
                "onu_id",
                "rx_power",
                "tx_power",
                "status",
            ]

            for field in preferred_fields:

                if field in all_fields:
                    field_list.append(field)
                    all_fields.remove(field)

            # Add any remaining API fields
            field_list.extend(
                sorted(all_fields)
            )

            # =====================================================
            # 8. Excel header
            # =====================================================

            for column_number, field in enumerate(
                field_list,
                start=1
            ):

                cell = worksheet.cell(
                    row=1,
                    column=column_number
                )

                cell.value = field

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center"
                )

            # =====================================================
            # 9. Add missing ONT records
            # =====================================================

            for row_number, item in enumerate(
                missing_onts,
                start=2
            ):

                for column_number, field in enumerate(
                    field_list,
                    start=1
                ):

                    value = item.get(field)

                    # Convert lists/dictionaries to string
                    # so Excel can store them.
                    if isinstance(value, (list, dict)):
                        value = str(value)

                    worksheet.cell(
                        row=row_number,
                        column=column_number
                    ).value = value

            # =====================================================
            # 10. Add summary sheet
            # =====================================================

            summary = workbook.create_sheet(
                title="Summary"
            )

            summary_data = [
                ("Signal API URL", API_URL),
                ("Total API records", len(data)),
                ("Unique MIB ONTs", len(api_serials)),
                ("ONTs in Django", len(customer_serials)),
                (
                    "MIB ONTs found in Django",
                    len(api_serials & customer_serials)
                ),
                (
                    "MIB ONTs missing in Django",
                    len(missing_serials)
                ),
            ]

            for row_number, (label, value) in enumerate(
                summary_data,
                start=1
            ):

                summary.cell(
                    row=row_number,
                    column=1
                ).value = label

                summary.cell(
                    row=row_number,
                    column=1
                ).font = Font(
                    bold=True
                )

                summary.cell(
                    row=row_number,
                    column=2
                ).value = value

            # =====================================================
            # 11. Freeze header row
            # =====================================================

            worksheet.freeze_panes = "A2"

            # =====================================================
            # 12. Add filter to header
            # =====================================================

            if field_list:
                last_column = get_column_letter(
                    len(field_list)
                )

                worksheet.auto_filter.ref = (
                    f"A1:{last_column}{len(missing_onts) + 1}"
                )

            # =====================================================
            # 13. Adjust column widths
            # =====================================================

            for column_cells in worksheet.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    try:
                        cell_length = len(
                            str(cell.value)
                        )

                        if cell_length > max_length:
                            max_length = cell_length

                    except Exception:
                        pass

                # Keep columns from becoming excessively wide
                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    40
                )

            summary.column_dimensions["A"].width = 35
            summary.column_dimensions["B"].width = 60

            # =====================================================
            # 14. Save Excel file
            # =====================================================

            output_file = os.path.join(
                os.getcwd(),
                "missing_mib_onts.xlsx"
            )

            workbook.save(output_file)

            # =====================================================
            # 15. Print result
            # =====================================================

            self.stdout.write("")
            self.stdout.write("=" * 80)
            self.stdout.write("RESULT")
            self.stdout.write("=" * 80)

            self.stdout.write(
                f"Total MIB ONTs       : {len(api_serials)}"
            )

            self.stdout.write(
                f"Found in Django      : "
                f"{len(api_serials & customer_serials)}"
            )

            self.stdout.write(
                f"Missing in Django    : {len(missing_serials)}"
            )

            self.stdout.write("")
            self.stdout.write(
                f"✅ Excel file created:"
            )

            self.stdout.write(
                output_file
            )

            # =====================================================
            # 16. Print missing ONTs to terminal too
            # =====================================================

            if missing_onts:

                self.stdout.write("")
                self.stdout.write(
                    "❌ MISSING MIB ONTs:"
                )
                self.stdout.write(
                    "-" * 80
                )

                for item in missing_onts:

                    serial = item.get(
                        "serial_number",
                        "-"
                    )

                    self.stdout.write(
                        str(serial)
                    )

                self.stdout.write(
                    "-" * 80
                )

            else:

                self.stdout.write("")
                self.stdout.write(
                    "✅ All MIB ONTs are present in Django."
                )

            self.stdout.write("")
            self.stdout.write(
                "MIB ONT check completed successfully."
            )

        except requests.exceptions.ConnectionError as e:

            self.stderr.write(
                f"❌ Signal API connection error: {e}"
            )

        except requests.exceptions.Timeout as e:

            self.stderr.write(
                f"❌ Signal API timeout: {e}"
            )

        except requests.exceptions.HTTPError as e:

            self.stderr.write(
                f"❌ Signal API HTTP error: {e}"
            )

        except Exception as e:

            self.stderr.write(
                f"❌ Error while checking MIB ONTs: {e}"
            )
